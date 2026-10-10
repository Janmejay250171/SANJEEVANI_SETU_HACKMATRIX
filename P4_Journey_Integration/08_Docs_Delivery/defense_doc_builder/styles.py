from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Palette Constants
COLOR_PRIMARY_HEX = "0B2545"      # Dark Midnight Navy
COLOR_SECONDARY_HEX = "134074"    # Deep Blue
COLOR_ACCENT_HEX = "0077B6"       # Vibrant Blue
COLOR_SUCCESS_HEX = "1B4332"      # Deep Forest Green
COLOR_WARNING_HEX = "854D0E"      # Rich Amber
COLOR_DANGER_HEX = "991B1B"       # Deep Red
COLOR_MUTED_HEX = "475569"        # Slate Gray
COLOR_BORDER_HEX = "CBD5E1"       # Light Gray Border
COLOR_BG_LIGHT_HEX = "F8FAFC"     # Off-white / Slate 50
COLOR_BG_ALERT_HEX = "FEF2F2"     # Light Red 50
COLOR_BG_WARN_HEX = "FFFBEB"      # Light Amber 50
COLOR_BG_SUCCESS_HEX = "F0FDF4"   # Light Green 50
COLOR_BG_INFO_HEX = "F0F9FF"      # Light Sky Blue 50
COLOR_BG_CODE_HEX = "F1F5F9"      # Code Block Background

COLOR_PRIMARY = RGBColor(11, 37, 69)
COLOR_SECONDARY = RGBColor(19, 64, 116)
COLOR_TEXT = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(71, 85, 105)
COLOR_DANGER = RGBColor(153, 27, 27)
COLOR_SUCCESS = RGBColor(27, 67, 50)
COLOR_WARNING = RGBColor(133, 77, 14)
COLOR_ACCENT = RGBColor(0, 119, 182)

def setup_document_styles(doc: Document):
    """Sets page margins and baseline typography styles."""
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(10)
    font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(3)

def add_header_footer(doc: Document):
    for section in doc.sections:
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SEHATSURE — COMPLETE TECHNICAL DEFENSE QUESTION BANK | FIN-01")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("CONFIDENTIAL & AUTHORITATIVE — Hackathon Final Round Defense & Technical Viva Masterbook")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8)
        frun.font.color.rgb = COLOR_MUTED

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def add_cover_page(doc: Document, title: str, subtitle: str, meta_rows: list):
    p_top_sp = doc.add_paragraph()
    p_top_sp.paragraph_format.space_before = Pt(40)

    # Decorative Top Banner Table
    banner_tbl = doc.add_table(rows=1, cols=1)
    banner_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    bcell = banner_tbl.cell(0, 0)
    set_cell_background(bcell, COLOR_PRIMARY_HEX)
    set_cell_margins(bcell, top=360, bottom=360, left=360, right=360)
    bp = bcell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brun_sup = bp.add_run("AI/ML HEALTHCARE INNOVATION HACKATHON 2026\nFINAL ROUND DEFENSE AUDIT & QUESTION BANK\n")
    brun_sup.font.name = "Calibri"
    brun_sup.font.size = Pt(12)
    brun_sup.font.bold = True
    brun_sup.font.color.rgb = RGBColor(186, 230, 253)

    brun = bp.add_run("🏥 SEHATSURE\n")
    brun.font.name = "Calibri"
    brun.font.size = Pt(32)
    brun.font.bold = True
    brun.font.color.rgb = RGBColor(255, 255, 255)

    brun_sub = bp.add_run("HOLISTIC OPTIMIZATION SYSTEM FOR POLICY-INTEGRATED ADMISSION & TREATMENT INTELLIGENCE")
    brun_sub.font.name = "Calibri"
    brun_sub.font.size = Pt(10.5)
    brun_sub.font.color.rgb = RGBColor(224, 242, 254)

    # Subtitle paragraph
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(28)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Table
    meta_tbl = doc.add_table(rows=len(meta_rows), cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (label, val) in enumerate(meta_rows):
        c0 = meta_tbl.cell(idx, 0)
        c1 = meta_tbl.cell(idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_background(c0, COLOR_BG_LIGHT_HEX)
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=90, bottom=90, left=130, right=130)
        set_cell_margins(c1, top=90, bottom=90, left=130, right=130)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_TEXT

    # Apply table borders
    for row in meta_tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                                f'<w:left w:val="none"/>'
                                f'<w:right w:val="none"/>'
                                f'</w:tcBorders>')
            tcPr.append(borders)

    # Bottom notice
    p_bot = doc.add_paragraph()
    p_bot.paragraph_format.space_before = Pt(36)
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_bot = p_bot.add_run(
        "⚠️ CRITICAL DEFENSE RULE:\n"
        "The actual source code and datasets are the sole source of truth.\n"
        "Never claim functionality not present in the codebase. Never fabricate capabilities."
    )
    r_bot.font.name = "Calibri"
    r_bot.font.size = Pt(9.5)
    r_bot.font.bold = True
    r_bot.font.color.rgb = COLOR_DANGER

    doc.add_page_break()

def add_heading_1(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_heading_2(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_paragraph(doc: Document, text: str, bold_prefix: str = "", italic: bool = False, color: RGBColor = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix + " ")
        rb.bold = True
        rb.font.name = 'Calibri'
        rb.font.size = Pt(10)
        rb.font.color.rgb = COLOR_PRIMARY
    rx = p.add_run(text)
    rx.font.name = 'Calibri'
    rx.font.size = Pt(10)
    rx.font.italic = italic
    rx.font.color.rgb = color or COLOR_TEXT
    return p

def add_bullet(doc: Document, title: str, text: str, bold_title: bool = True):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if title:
        rt = p.add_run(title + (": " if bold_title else " "))
        rt.font.name = 'Calibri'
        rt.font.bold = bold_title
        rt.font.color.rgb = COLOR_PRIMARY
    rx = p.add_run(text)
    rx.font.name = 'Calibri'
    rx.font.size = Pt(10)
    rx.font.color.rgb = COLOR_TEXT
    return p

def add_callout(doc: Document, text: str, title: str = "IMPORTANT NOTICE", alert_type: str = "info"):
    bg_map = {
        "info": (COLOR_BG_INFO_HEX, COLOR_SECONDARY_HEX, COLOR_SECONDARY),
        "warning": (COLOR_BG_WARN_HEX, COLOR_WARNING_HEX, COLOR_WARNING),
        "danger": (COLOR_BG_ALERT_HEX, COLOR_DANGER_HEX, COLOR_DANGER),
        "success": (COLOR_BG_SUCCESS_HEX, COLOR_SUCCESS_HEX, COLOR_SUCCESS)
    }
    bg_hex, border_hex, title_color = bg_map.get(alert_type, bg_map["info"])

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=130, bottom=130, left=160, right=160)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
                        f'<w:top w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'<w:bottom w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = title_color

    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_TEXT

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(3)

def add_code_block(doc: Document, code_str: str, language: str = ""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, COLOR_BG_CODE_HEX)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="16" w:space="0" w:color="{COLOR_SECONDARY_HEX}"/>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05

    if language:
        rl = p.add_run(f"// {language}\n")
        rl.font.name = 'Consolas'
        rl.font.size = Pt(8.5)
        rl.font.color.rgb = COLOR_MUTED
        rl.font.italic = True

    rc = p.add_run(code_str)
    rc.font.name = 'Consolas'
    rc.font.size = Pt(8.5)
    rc.font.color.rgb = RGBColor(15, 23, 42)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(3)

def add_qa_card(doc: Document, q_num: str, question: str, ideal_answer: str, evidence: str, followup: str = "", status_tag: str = ""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, COLOR_BG_LIGHT_HEX)
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

    tcPr = cell._tc.get_or_add_tcPr()
    border_color = COLOR_SECONDARY_HEX
    if status_tag == "CRITICAL" or status_tag == "DANGER":
        border_color = COLOR_DANGER_HEX
    elif status_tag == "WARNING":
        border_color = COLOR_WARNING_HEX
    elif status_tag == "SUCCESS":
        border_color = COLOR_SUCCESS_HEX

    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:left w:val="single" w:sz="20" w:space="0" w:color="{border_color}"/>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

    # Question line
    p_q = cell.paragraphs[0]
    p_q.paragraph_format.space_before = Pt(2)
    p_q.paragraph_format.space_after = Pt(3)
    p_q.paragraph_format.line_spacing = 1.15
    rq_num = p_q.add_run(f"Q{q_num}: ")
    rq_num.bold = True
    rq_num.font.name = "Calibri"
    rq_num.font.size = Pt(10.5)
    rq_num.font.color.rgb = COLOR_PRIMARY

    rq = p_q.add_run(question)
    rq.bold = True
    rq.font.name = "Calibri"
    rq.font.size = Pt(10.5)
    rq.font.color.rgb = COLOR_PRIMARY

    if status_tag:
        tag_colors = {
            "CRITICAL": COLOR_DANGER,
            "HIGH": COLOR_DANGER,
            "MEDIUM": COLOR_WARNING,
            "LOW": COLOR_MUTED,
            "SAFE": COLOR_SUCCESS,
            "QUALIFY": COLOR_WARNING
        }
        rtag = p_q.add_run(f"  [{status_tag}]")
        rtag.bold = True
        rtag.font.name = "Calibri"
        rtag.font.size = Pt(8.5)
        rtag.font.color.rgb = tag_colors.get(status_tag, COLOR_MUTED)

    # Ideal Answer
    p_a = cell.add_paragraph()
    p_a.paragraph_format.space_before = Pt(2)
    p_a.paragraph_format.space_after = Pt(2)
    p_a.paragraph_format.line_spacing = 1.15
    ra_lbl = p_a.add_run("Ideal Answer: ")
    ra_lbl.bold = True
    ra_lbl.font.name = "Calibri"
    ra_lbl.font.size = Pt(9.5)
    ra_lbl.font.color.rgb = COLOR_SECONDARY
    ra_txt = p_a.add_run(ideal_answer)
    ra_txt.font.name = "Calibri"
    ra_txt.font.size = Pt(9.5)
    ra_txt.font.color.rgb = COLOR_TEXT

    # Evidence line
    if evidence:
        p_e = cell.add_paragraph()
        p_e.paragraph_format.space_before = Pt(1)
        p_e.paragraph_format.space_after = Pt(2)
        p_e.paragraph_format.line_spacing = 1.15
        re_lbl = p_e.add_run("Evidence in Code/Files: ")
        re_lbl.bold = True
        re_lbl.font.name = "Calibri"
        re_lbl.font.size = Pt(9)
        re_lbl.font.color.rgb = COLOR_MUTED
        re_txt = p_e.add_run(evidence)
        re_txt.font.name = "Calibri"
        re_txt.font.size = Pt(9)
        re_txt.font.italic = True
        re_txt.font.color.rgb = COLOR_TEXT

    # Followup line
    if followup:
        p_f = cell.add_paragraph()
        p_f.paragraph_format.space_before = Pt(1)
        p_f.paragraph_format.space_after = Pt(2)
        p_f.paragraph_format.line_spacing = 1.15
        rf_lbl = p_f.add_run("Potential Follow-up: ")
        rf_lbl.bold = True
        rf_lbl.font.name = "Calibri"
        rf_lbl.font.size = Pt(9)
        rf_lbl.font.color.rgb = COLOR_WARNING
        rf_txt = p_f.add_run(followup)
        rf_txt.font.name = "Calibri"
        rf_txt.font.size = Pt(9)
        rf_txt.font.color.rgb = COLOR_TEXT

    # Space after card
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(3)

def format_table(tbl, col_widths, headers, rows):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    for i, h_text in enumerate(headers):
        cell = hdr_cells[i]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.add_row().cells
        bg_color = COLOR_BG_LIGHT_HEX if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row_cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(str(cell_value))
            run.font.name = 'Calibri'
            run.font.size = Pt(8.5)
            run.font.color.rgb = COLOR_TEXT

    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                                f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                                f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER_HEX}"/>'
                                f'</w:tcBorders>')
            tcPr.append(borders)
