#!/usr/bin/env python3
"""
build_defense_question_bank.py
------------------------------
Generates SEHATSURE_COMPLETE_TECHNICAL_DEFENSE_QUESTION_BANK.docx
for the AI/ML Healthcare Hackathon 2026 (HackMatrix FIN-01).

Assembles:
- Part 1: Cover Page, How to Use, Architecture Overview, Inventory, FIN-01 Alignment, Sections A, B, C
- Part 2: Section D (File-by-File), Section E (Code Questions), Section F (Code Snippet Defense)
- Part 3: Section G (Policy Extraction), Section H (AI/ML), Section I (Cost Engine Detailed), Section J (Numerical Viva), Section K (Policy Semantics)
- Part 4: Section L (Hospital / Network Logic), Section M (Scoring / Ranking), Section N (Database/Datasets), Sections O, P, Q, R, S, T, U, V, W, X
- Part 5: Section Y (Judge Attacks), Section Z (Code Viva Checklist), Section AA (Formula Viva), Sections AB, AC, AD, AE, AF, AG
- Part 6: Section AH (32 30s Answers), Section AI (11 1m Pitches), Section AJ (105 Rapid-Fire Q&As), Section AK (105 Deep Technical Q&As), Sections AL, AM, AN, AO
"""

import os
import sys
import docx
from docx import Document
from defense_doc_builder.styles import setup_document_styles, add_header_footer
from defense_doc_builder.part1_foundation import build_part1
from defense_doc_builder.part2_files_and_code import build_part2
from defense_doc_builder.part3_ai_and_math import build_part3
from defense_doc_builder.part4_hospital_system import build_part4
from defense_doc_builder.part5_attacks_and_audits import build_part5
from defense_doc_builder.part6_master_qa_and_checklists import build_part6

def main():
    print("=" * 70)
    print("SEHATSURE — GENERATING COMPLETE TECHNICAL DEFENSE QUESTION BANK")
    print("National AI/ML Healthcare Hackathon 2026 | Problem Statement FIN-01")
    print("=" * 70)

    doc = Document()
    setup_document_styles(doc)

    print("\n[1/6] Assembling Part 1: Foundation, Inventory, FIN-01, Fundamentals & Architecture...")
    build_part1(doc)

    print("\n[2/6] Assembling Part 2: File-by-File Audit, Code Questions & Code Snippet Defense...")
    build_part2(doc)

    print("\n[3/6] Assembling Part 3: AI/ML, Policy Extraction, Detailed Cost Engine & Numerical Viva...")
    build_part3(doc)

    print("\n[4/6] Assembling Part 4: Network Logic, Scoring, Database, Security, API & Edge Cases...")
    build_part4(doc)

    print("\n[5/6] Assembling Part 5: Judge Attacks, Code Viva, Formula Sheet, Claim Audit & Role Defense...")
    build_part5(doc)

    print("\n[6/6] Assembling Part 6: 30s Answers, 1m Pitches, 100+ Rapid-Fire, 100+ Deep Tech & Checklists...")
    build_part6(doc)

    add_header_footer(doc)

    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "SEHATSURE_COMPLETE_TECHNICAL_DEFENSE_QUESTION_BANK.docx"))
    doc.save(output_path)

    file_size_kb = os.path.getsize(output_path) / 1024
    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)

    print("\n" + "=" * 70)
    print("SUCCESS: Comprehensive Defense Question Bank Generated Successfully!")
    print(f" -> Output Path: {output_path}")
    print(f" -> File Size:   {file_size_kb:.2f} KB")
    print(f" -> Paragraphs:  {para_count}")
    print(f" -> Tables:      {table_count}")
    print("=" * 70)

if __name__ == "__main__":
    main()
